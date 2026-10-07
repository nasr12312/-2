"""Read-only reconciliation of school workbooks and a private database snapshot."""
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r'C:\Users\abd22\OneDrive\Desktop\مدرسة غراس الأخلاق')

def text(value):
    return re.sub(r'\s+', ' ', str(value or '').strip())

def normalized(value):
    return re.sub(r'[\sـ\u064b-\u065f\u0670]', '', unicodedata.normalize('NFKD', text(value))).translate(str.maketrans('أإآٱى', 'ااااي'))

def number(value):
    return text(value).removesuffix('.0').translate(str.maketrans('٠١٢٣٤٥٦٧٨٩', '0123456789'))

def source_rows(filename, program):
    workbook = load_workbook(SOURCE / filename, read_only=True, data_only=True)
    result = []
    for sheet in workbook:
        iterator = sheet.iter_rows(values_only=True)
        headers = [text(x) for x in next(iterator)]
        columns = {name: headers.index(name) for name in ['اسم الطالب', 'اسم المرحلة', 'اسم الصف']}
        if program == 'bilingual':
            columns.update({name: headers.index(name) for name in ['رقم الهوية', 'اسم الفصل']})
        for line, row in enumerate(iterator, 2):
            name = text(row[columns['اسم الطالب']])
            if len(name.split()) < 2:
                continue
            section = number(row[columns['اسم الفصل']]) if program == 'bilingual' else ''
            if not section or section == 'قائمة المسجلين الجدد':
                section = 'غير موزع'
            item = dict(name=name, stage=text(row[columns['اسم المرحلة']]), grade=text(row[columns['اسم الصف']]),
                        section=section, program=program, sheet=sheet.title, row=line,
                        national_id=number(row[columns['رقم الهوية']]) if program == 'bilingual' else '')
            if program == 'bilingual' and not re.fullmatch(r'\d{10}', item['national_id']):
                raise ValueError(f'Invalid student identifier at {sheet.title}:{line}')
            result.append(item)
    workbook.close()
    return result

snapshot = json.loads((ROOT / 'private-data/roster-database-oct7.json').read_text(encoding='utf-8-sig'))
bilingual = source_rows('طلاب ثنائي اللغة.xlsx', 'bilingual')
diploma = source_rows('بيانات طلاب الدبلوما بتاريخ 27-8-2026 (1).xlsx', 'diploma')
by_name = defaultdict(list)
for student in bilingual:
    by_name[(normalized(student['name']), normalized(student['stage']), normalized(student['grade']))].append(student)
unmatched = []
for student in diploma:
    matches = by_name[(normalized(student['name']), normalized(student['stage']), normalized(student['grade']))]
    ids = {m['national_id'] for m in matches}
    if len(ids) == 1:
        student['national_id'] = next(iter(ids))
    else:
        unmatched.append(student)

classes = {c['id']: c for c in snapshot['classes']}
source = bilingual + diploma
expected = {(s['program'], s['national_id']): s for s in source if s['national_id']}
actual = {(s['program'], s['national_id']): s for s in snapshot['students'] if s['active']}
differences = []
for key, student in expected.items():
    saved = actual.get(key)
    if not saved:
        differences.append({'issue': 'missing_student', 'source': student})
        continue
    classroom = classes[saved['class_id']]
    if normalized(student['name']) != normalized(saved['name']):
        differences.append({'issue': 'name_mismatch', 'source': student, 'database': saved})
    for field in ['stage', 'grade', 'section']:
        if normalized(student[field]) != normalized(classroom[field]):
            differences.append({'issue': field + '_mismatch', 'source': student, 'database': saved, 'class': classroom})
for key in actual.keys() - expected.keys():
    differences.append({'issue': 'not_in_source', 'database': actual[key]})
source_counts = Counter((s['program'], s['national_id']) for s in source if s['national_id'])
database_counts = Counter((s['program'], s['national_id']) for s in snapshot['students'] if s['active'])
class_counts = Counter((c['program_id'], normalized(c['stage']), normalized(c['grade']), normalized(c['section'])) for c in snapshot['classes'] if c['active'])
source_classes = {(s['program'], normalized(s['stage']), normalized(s['grade']), normalized(s['section'])) for s in source}
database_classes = set(class_counts)
summary = {
    'source_bilingual_rows': len(bilingual), 'source_diploma_rows': len(diploma),
    'unique_children': len({s['national_id'] for s in source if s['national_id']}),
    'children_in_both_programs': len({s['national_id'] for s in bilingual} & {s['national_id'] for s in diploma}),
    'duplicate_source_enrollments': sum(n-1 for n in source_counts.values() if n > 1),
    'duplicate_database_enrollments': sum(n-1 for n in database_counts.values() if n > 1),
    'duplicate_active_classes': sum(n-1 for n in class_counts.values() if n > 1),
    'source_class_count': len(source_classes), 'database_active_class_count': len(database_classes),
    'unmatched_diploma_rows': len(unmatched), 'differences': len(differences),
    'missing_classes': len(source_classes-database_classes), 'extra_classes': len(database_classes-source_classes),
    'difference_types': dict(Counter(d['issue'] for d in differences))
}
(ROOT / 'private-data/roster-audit-oct7.json').write_text(json.dumps({'summary': summary, 'differences': differences, 'unmatched_diploma': unmatched, 'source': source}, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(summary, ensure_ascii=False))
