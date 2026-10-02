from __future__ import annotations

import csv
import hashlib
import json
import re
import secrets
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


PROJECT = Path(__file__).resolve().parents[1]
SOURCE = Path(r"C:\Users\abd22\OneDrive\Desktop\مدرسة غراس الأخلاق")
OUTPUT = PROJECT / "private-data"

DIPLOMA_FILE = SOURCE / "بيانات طلاب الدبلوما بتاريخ 27-8-2026 (1).xlsx"
BILINGUAL_FILE = SOURCE / "طلاب ثنائي اللغة.xlsx"

FORM_FILES = [
    "استمارة متابعة قرآن أول (1).docx",
    "__استمارة القرآن الكريم - ثاني تحفيظ - الفصل الأول 1448هـ_ (2).docx",
    "استمارة القرآن الكريم - ثالث تحفيظ - تعبئة الخميس - نهائي.docx",
    "استمارة القرآن رابع تحفيظ الفصل الأول 1448هـ.docx",
    "استمارة القرآن خامس فصل دراسي أول تحفيظ  1448هـ.docx",
    "__استمارة القرآن الكريم - أول متوسط دبلومة تحفيظ - الفصل الأول 1448هـ_.docx",
    "__استمارة القرآن الكريم - ثاني متوسط دبلومة تحفيظ - الفصل الأول 1448هـ_.docx",
    "__استمارة القرآن الكريم - ثالث متوسط تحفيظ - الفصل الأول 1448هـ_.docx",
]

TEACHERS = [
    ("عبد الرحمن علي نصر الله", "head_teacher"),
    ("أحمد عبد الله العامري", "teacher"),
    ("ياسين حسين علي", "teacher"),
    ("هاشم الهنداوي", "supervisor_teacher"),
    ("باسم مرزوق", "supervisor"),
    ("أحمد مصطفى محمود", "teacher"),
    ("أحمد الرفاعي", "teacher"),
]


def clean(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return re.sub(r"\s+", " ", text)


def digits(value) -> str:
    return re.sub(r"\D", "", clean(value).translate(str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789")))


def stable_id(prefix: str, *parts: str) -> str:
    value = "|".join(clean(x) for x in parts)
    return f"{prefix}_{hashlib.sha256(value.encode('utf-8')).hexdigest()[:16]}"


def compact_header(value) -> str:
    return clean(value).replace(" ", "").replace("ـ", "")


def header_index(row, candidates: list[str]) -> int | None:
    normalized = [compact_header(v) for v in row]
    for i, value in enumerate(normalized):
        if any(compact_header(c) in value for c in candidates):
            return i
    return None


def class_record(program: str, stage: str, grade: str, section: str) -> dict:
    stage = clean(stage) or "غير محدد"
    grade = clean(grade) or "غير محدد"
    section = clean(section) or "أ"
    return {
        "id": stable_id("class", program, stage, grade, section),
        "program": program,
        "stage": stage,
        "grade": grade,
        "section": section,
        "name": f"{grade} - {section}",
    }


def read_diploma() -> tuple[list[dict], dict[str, dict]]:
    workbook = load_workbook(DIPLOMA_FILE, data_only=True, read_only=True)
    students: list[dict] = []
    classes: dict[str, dict] = {}
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        header_row = next((i for i, row in enumerate(rows[:12]) if any("اسم" in clean(v) for v in row)), 0)
        headers = rows[header_row]
        name_i = header_index(headers, ["اسمالطالب", "الاسم"])
        stage_i = header_index(headers, ["المرحلة"])
        grade_i = header_index(headers, ["الصف"])
        if name_i is None:
            continue
        for row in rows[header_row + 1 :]:
            name = clean(row[name_i] if name_i < len(row) else "")
            if len(name.split()) < 2 or "اسم الطالب" in name or name.isdigit():
                continue
            stage = clean(row[stage_i]) if stage_i is not None and stage_i < len(row) else ""
            grade = clean(row[grade_i]) if grade_i is not None and grade_i < len(row) else clean(sheet.title)
            klass = class_record("دبلومة", stage, grade, "أ")
            classes[klass["id"]] = klass
            students.append({
                "id": stable_id("student", "دبلومة", name, stage, grade),
                "name": name,
                "national_id": "",
                "program": "دبلومة",
                "class_id": klass["id"],
                "account_status": "waiting_national_id",
            })
    return students, classes


def read_bilingual() -> tuple[list[dict], dict[str, dict]]:
    workbook = load_workbook(BILINGUAL_FILE, data_only=True, read_only=True)
    students: list[dict] = []
    classes: dict[str, dict] = {}
    seen: set[tuple[str, str]] = set()
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        header_row = next((i for i, row in enumerate(rows[:15]) if sum(bool(clean(v)) for v in row) >= 4 and any("اسم" in clean(v) for v in row)), None)
        if header_row is None:
            continue
        headers = rows[header_row]
        name_i = header_index(headers, ["اسمالطالب", "الطالب"])
        nid_i = header_index(headers, ["الرقمالوطني", "رقمالهوية", "هويةالطالب"])
        stage_i = header_index(headers, ["المرحلة"])
        grade_i = header_index(headers, ["الصف"])
        section_i = header_index(headers, ["الشعبة", "الفصل"])
        if name_i is None or nid_i is None:
            continue
        for row in rows[header_row + 1 :]:
            name = clean(row[name_i] if name_i < len(row) else "")
            national_id = digits(row[nid_i] if nid_i < len(row) else "")
            if not name or not national_id or len(national_id) < 6:
                continue
            dedupe_key = (name, national_id)
            if dedupe_key in seen:
                continue
            seen.add(dedupe_key)
            stage = clean(row[stage_i]) if stage_i is not None and stage_i < len(row) else ""
            grade = clean(row[grade_i]) if grade_i is not None and grade_i < len(row) else clean(sheet.title)
            section = clean(row[section_i]) if section_i is not None and section_i < len(row) else "غير موزع"
            if section == "قائمة المسجلين الجدد":
                section = "غير موزع"
            klass = class_record("ثنائي اللغة", stage, grade, section)
            classes[klass["id"]] = klass
            students.append({
                "id": stable_id("student", "ثنائي اللغة", national_id),
                "name": name,
                "national_id": national_id,
                "program": "ثنائي اللغة",
                "class_id": klass["id"],
                "account_status": "active",
            })
    return students, classes


def form_grade(filename: str) -> tuple[str, str]:
    normalized = filename.replace("أ", "أ").replace("آ", "آ")
    if "أول متوسط" in normalized:
        return "المتوسطة", "الأول المتوسط"
    if "ثاني متوسط" in normalized:
        return "المتوسطة", "الثاني المتوسط"
    if "ثالث متوسط" in normalized:
        return "المتوسطة", "الثالث المتوسط"
    if "ثاني" in normalized:
        return "الابتدائية", "الثاني الابتدائي"
    if "ثالث" in normalized:
        return "الابتدائية", "الثالث الابتدائي"
    if "رابع" in normalized:
        return "الابتدائية", "الرابع الابتدائي"
    if "خامس" in normalized:
        return "الابتدائية", "الخامس الابتدائي"
    return "الابتدائية", "الأول الابتدائي"


def read_plan_forms() -> list[dict]:
    entries: list[dict] = []
    for filename in FORM_FILES:
        path = SOURCE / filename
        document = Document(path)
        stage, grade = form_grade(filename)
        for table in document.tables:
            for row in table.rows:
                values = [clean(cell.text) for cell in row.cells]
                if len(values) < 5:
                    continue
                week = digits(values[0])
                row_kind = values[1]
                date_or_day = values[2]
                surah = values[3]
                verses = values[4]
                if not week or not row_kind or not surah:
                    continue
                if "الغيب" in row_kind or "حفظ" in row_kind:
                    assignment_type = "حفظ"
                elif "التلاوة" in row_kind or "تلاوة" in row_kind:
                    assignment_type = "تلاوة"
                else:
                    continue
                assignment_text = " — ".join(dict.fromkeys(x for x in [surah, verses] if x))
                entries.append({
                    "id": stable_id("plan", filename, week, date_or_day, assignment_type, assignment_text),
                    "program": "دبلومة",
                    "stage": stage,
                    "grade": grade,
                    "week_label": week,
                    "day_name": date_or_day,
                    "assignment_type": assignment_type,
                    "assignment_text": assignment_text,
                    "source_form": filename,
                })
    unique = {entry["id"]: entry for entry in entries}
    return list(unique.values())


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    diploma_students, diploma_classes = read_diploma()
    bilingual_students, bilingual_classes = read_bilingual()
    # Match the same child across programs only by exact name, stage and grade.
    for student in diploma_students:
        classroom = diploma_classes[student['class_id']]
        matches = [s for s in bilingual_students if s['name'] == student['name']
                   and bilingual_classes[s['class_id']]['stage'] == classroom['stage']
                   and bilingual_classes[s['class_id']]['grade'] == classroom['grade']]
        if len(matches) == 1:
            student['national_id'] = matches[0]['national_id']
            student['account_status'] = 'active'
            section = 'غير موزع'
        else:
            section = 'غير موزع'
        klass = class_record('دبلومة', classroom['stage'], classroom['grade'], section)
        diploma_classes[klass['id']] = klass
        student['class_id'] = klass['id']
    diploma_classes = {k:v for k,v in diploma_classes.items() if any(s['class_id']==k for s in diploma_students)}
    students = diploma_students + bilingual_students
    classes = {**diploma_classes, **bilingual_classes}
    plans = read_plan_forms()
    plans += [{**entry, 'id': 'shared_' + entry['id'], 'program': 'ثنائي اللغة'} for entry in list(plans)]

    codes_path = OUTPUT / 'teacher-access.json'
    teacher_codes = json.loads(codes_path.read_text(encoding='utf-8')) if codes_path.exists() else {}
    for name, _ in TEACHERS:
        teacher_codes.setdefault(name, str(7000000000 + secrets.randbelow(999999999)))
    codes_path.write_text(json.dumps(teacher_codes, ensure_ascii=False), encoding='utf-8')
    teachers = [
        {"id": stable_id("teacher", name), "access_code": teacher_codes[name], "name": name, "role": role}
        for name, role in TEACHERS
    ]
    payload = {
        "version": 1,
        "programs": [{"id": "diploma", "name": "دبلومة"}, {"id": "bilingual", "name": "ثنائي اللغة"}],
        "classes": sorted(classes.values(), key=lambda x: (x["program"], x["stage"], x["grade"], x["section"])),
        "students": students,
        "teachers": teachers,
        "quran_plan_entries": plans,
    }
    (OUTPUT / "school-import.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    with (OUTPUT / "login-codes.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["النوع", "الاسم", "البرنامج", "رمز الدخول", "الحالة"])
        for teacher in teachers:
            writer.writerow(["معلم", teacher["name"], "", teacher["access_code"], "نشط"])
        for student in students:
            writer.writerow([
                "ولي أمر",
                student["name"],
                student["program"],
                student["national_id"],
                "نشط" if student["national_id"] else "بانتظار الرقم الوطني",
            ])

    ids = [s["national_id"] for s in bilingual_students if s["national_id"]]
    summary = {
        "student_count": len(students),
        "unique_student_count": len({s['national_id'] for s in students if s['national_id']}),
        "numbered_class_count": sum(x['section'].isdigit() for x in classes.values()),
        "pending_distribution_groups": sum(x['section']=='غير موزع' for x in classes.values()),
        "diploma_student_count": len(diploma_students),
        "bilingual_student_count": len(bilingual_students),
        "active_parent_accounts": len(ids),
        "waiting_for_national_id": sum(not s["national_id"] for s in students),
        "duplicate_national_ids": len(ids) - len(set(ids)),
        "class_count": len(classes),
        "teacher_count": len(teachers),
        "plan_entry_count": len(plans),
        "stored_identity_fields": ["name", "national_id"],
    }
    (OUTPUT / "import-summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False))


if __name__ == "__main__":
    main()
